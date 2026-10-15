"""
"Preparar factura": los datos de la factura de un pedido y la FACTURA EN EXCEL, completa, como la
plantilla de factura de Yikon (15/10/2026: antes era un Excel básico y Aina la montaba a mano).

  - Logo; Seller (Ajustes → Datos de Yikon, app/datos_empresa.py), Billing address (los datos de
    facturación del cliente) y Delivery address (sus datos de envío); fecha, nº de factura (= la
    referencia interna del pedido) y su nº de pedido.
  - EORI: si el país del cliente es de la UE (app/paises.UE), el de Yikon en Seller y el del cliente
    en su dirección (si los tienen).
  - Los productos, con las columnas del FORMATO del cliente (ficha del cliente → Facturación):
    A kits sin precio, B kits con precio por línea, C lista con precio unitario; y el HS Code de
    cada producto.
  - Los totales según su régimen de IVA y sus gastos de envío pactados (con fórmulas, para que
    se recalculen si Aina toca algo); el uso (RUO / IVD).
  - Transferencia (banco, IBAN, SWIFT de Ajustes), plazo de entrega, Incoterm, transportista,
    «Confirmed & Accepted By» y el sello de Yikon.
La pantalla va en el idioma de la web; el Excel, en inglés (como las facturas). Aina la revisa,
la pasa a PDF y la registra como siempre (Facturas → registrar).
"""
import os
from datetime import date
from decimal import Decimal

from app.inventario.reservas import reservas_de


from flask_babel import lazy_gettext as _l

def _lotes(reservas, clave, usadas):
    """Los lotes de esa clave, solo la primera vez que sale (un mismo producto puede ir suelto
    y en un kit: sus lotes se enseñan una vez)."""
    if clave in usadas:
        return []
    usadas.add(clave)
    return reservas.get(clave, [])


def filas_productos(pedido) -> list:
    """Una fila por producto: [{"producto", "unidad", "referencia", "item", "cantidad", "unitario", "precio",
    "lotes": [(lote, unidades)], "primera", "grupo"}]. "producto" y "unidad" (el kit y cuántos) solo en
    la primera fila de cada kit; "grupo" = cuántas filas ocupa.

    Precios (formatos B y C): los productos de un kit no tienen precio propio en el pedido (lo paga el
    kit), así que llevan el de catálogo (avisos_kits() dice si no suman el precio del kit). Los sueltos,
    el del pedido. "precio" es el de la línea SIN descuento: el descuento va en los totales."""
    reservas, usadas, filas = reservas_de(pedido), set(), []
    for k in pedido.kits:
        lotes_kit = _lotes(reservas, ("kit", k.kit_id), usadas) if k.kit.cerrado else []
        if not k.lineas:          # un kit sin productos (raro): al menos su fila, con su precio
            filas.append({"producto": k.kit.nombre, "unidad": k.cantidad, "referencia": k.kit.referencia,
                          "item": k.kit.nombre, "cantidad": k.cantidad, "unitario": Decimal(k.precio_unitario or 0),
                          "precio": k.subtotal, "lotes": lotes_kit, "primera": True, "grupo": 1, "hs": ""})
        for i, l in enumerate(k.lineas):
            lotes = lotes_kit if i == 0 else []
            if not k.kit.cerrado:
                lotes = _lotes(reservas, ("producto", l.producto_id), usadas)
            unitario = Decimal(l.producto.precio or 0)
            filas.append({"producto": k.kit.nombre if i == 0 else "", "unidad": k.cantidad if i == 0 else "",
                          "referencia": l.producto.referencia, "item": l.producto.nombre, "cantidad": l.cantidad,
                          "unitario": unitario, "precio": (unitario * l.cantidad).quantize(Decimal("0.01")),
                          "lotes": lotes, "primera": i == 0, "grupo": len(k.lineas), "hs": l.producto.codigo_hs or ""})
    for l in pedido.lineas_sueltas:
        unitario = Decimal(l.precio_unitario or 0)
        filas.append({"producto": l.producto.nombre, "unidad": l.cantidad, "referencia": l.producto.referencia,
                      "item": l.producto.nombre, "cantidad": l.cantidad, "suelto": True, "unitario": unitario,
                      "precio": (unitario * l.cantidad).quantize(Decimal("0.01")),
                      "lotes": _lotes(reservas, ("producto", l.producto_id), usadas), "primera": True, "grupo": 1,
                      "hs": l.producto.codigo_hs or ""})
    return filas


def avisos_kits(pedido) -> list:
    """Kits cuyos productos, a precio de catálogo, no suman lo que cuesta el kit en el pedido:
    [(nombre del kit, suma de catálogo, precio del kit en el pedido)]."""
    avisos = []
    for k in pedido.kits:
        if not k.lineas:
            continue
        suma = sum((Decimal(l.producto.precio or 0) * l.cantidad for l in k.lineas), Decimal("0")).quantize(Decimal("0.01"))
        if suma != k.subtotal:
            avisos.append((k.kit.nombre, suma, k.subtotal))
    return avisos


IVA_ESPANA = Decimal("21")


def totales(pedido, regimen="es", envio=None) -> dict:
    """Precio total sin descuento, % de descuento (si hay descuentos por línea, el equivalente
    sobre el total), subtotal, envío, % de IVA (None si hay varios tipos), IVA y total.
    Fuera de España (UE o exportación) no hay IVA. En España el envío también lleva IVA."""
    bruto = (sum((k.subtotal for k in pedido.kits), Decimal("0"))
             + sum((Decimal(l.cantidad or 0) * Decimal(l.precio_unitario or 0) for l in pedido.lineas_sueltas),
                   Decimal("0"))).quantize(Decimal("0.01"))
    neto = pedido.base_imponible
    descuento = ((1 - neto / bruto) * 100).quantize(Decimal("0.01")) if bruto else Decimal("0")
    envio = Decimal(envio or 0).quantize(Decimal("0.01"))
    tipos = {Decimal(x.iva_pct) for x in [*pedido.kits, *pedido.lineas_sueltas] if x.subtotal}
    if not tipos:          # (todo a 0 €: el tipo de sus líneas, no «mixed»)
        tipos = {Decimal(x.iva_pct) for x in [*pedido.kits, *pedido.lineas_sueltas]}
    iva_pct = tipos.pop() if len(tipos) == 1 else None
    con_iva = regimen == "es"
    iva = Decimal("0.00")
    if con_iva:
        iva = (pedido.total_iva + envio * (iva_pct or IVA_ESPANA) / 100).quantize(Decimal("0.01"))
    return {"bruto": bruto, "descuento": descuento, "neto": neto, "envio": envio, "con_iva": con_iva,
            "iva_pct": iva_pct, "iva": iva, "total": neto + envio + iva}


def lista_totales(t) -> list:
    """Las filas de totales que tocan, en el orden de las facturas: [(clave, valor)]. El descuento
    solo si hay, el envío si se cobra y el IVA si es en España. iva_pct None = varios tipos."""
    filas = [("bruto", t["bruto"])]
    if t["descuento"]:
        filas += [("descuento", t["descuento"]), ("neto", t["neto"])]
    if t["envio"]:
        filas.append(("envio", t["envio"]))
    if t["con_iva"]:
        filas += [("iva_pct", t["iva_pct"]), ("iva", t["iva"])]
    filas.append(("total", t["total"]))
    return filas


# Nombres de los totales: en el Excel, como en las facturas (inglés); en pantalla, traducidos
TOTALES_EXCEL = {"bruto": "Total price", "descuento": "Discount [%]", "neto": "Subtotal less discount",
                 "envio": "Shipment", "iva_pct": "% IVA", "iva": "Total tax", "total": "Total amount"}
TOTALES_PANTALLA = {"bruto": _l("Precio total"), "descuento": _l("Descuento"), "neto": _l("Subtotal con descuento"),
                    "envio": _l("Envío"), "iva_pct": _l("IVA"), "iva": _l("Importe del IVA"), "total": _l("Total")}

# Columnas de cada formato: en el Excel como en las facturas; en pantalla, traducidas.
# Claves de cada fila de filas_productos(). En A y B, "Product" ocupa tres columnas (A:C).
COLUMNAS = {
    "A": ["producto", "unidad", "referencia", "item", "cantidad", "hs"],
    "B": ["producto", "unidad", "referencia", "item", "cantidad", "hs", "precio"],
    "C": ["item", "referencia", "hs", "unitario", "cantidad", "precio"],
}
COLUMNAS_EXCEL = {"producto": "PRODUCT", "unidad": "Unit", "referencia": "Catalog No.", "item": "Item",
                  "cantidad": "Qty.", "unitario": "Unit Price", "precio": "Price", "hs": "HS Code"}
COLUMNAS_PANTALLA = {"producto": _l("Producto"), "unidad": _l("Unid."), "referencia": _l("Referencia"),
                     "item": _l("Artículo"), "cantidad": _l("Cant."), "unitario": _l("Precio unitario"),
                     "precio": _l("Precio"), "hs": _l("HS Code")}
# El uso del pedido (lo elige el cliente al pedir), como sale en la factura
USO_EXCEL = {"ruo": "FOR RESEARCH USE ONLY (RUO)", "ivd": "FOR IN VITRO DIAGNOSTIC USE (IVD)"}


def _lineas_lugar(d):
    """Dirección, «CP Ciudad, Provincia» (la provincia en la misma línea y solo si no es la ciudad:
    «03001 Alicante», no «Alicante, Alicante») y país, una por renglón (sin las vacías)."""
    ciudad, provincia = d.get("ciudad") or "", d.get("provincia") or ""
    lugar = " ".join(y for y in (d.get("codigo_postal"), ciudad) if y)
    if provincia and provincia.strip().lower() != ciudad.strip().lower():
        lugar = ", ".join(y for y in (lugar, provincia) if y)
    return [x for x in (d.get("direccion"), lugar, (d.get("pais") or "").upper()) if x]


def datos_para_facturar(pedido, hoy=None, envio=None) -> dict:
    """Todo lo de "Preparar factura" con el formato, el IVA y el envío del cliente.
    envio: otro importe de envío solo para esta factura (None = el pactado con el cliente)."""
    from flask_babel import force_locale

    from app import paises
    hoy = hoy or date.today()
    c = pedido.cliente
    formato = c.formato_factura if c.formato_factura in COLUMNAS else "A"
    regimen = c.regimen_iva or "es"
    envio_pactado = c.gastos_envio
    if envio is None:
        envio = envio_pactado
    ue = paises.es_ue(c.pais)
    with force_locale("en"):                     # el Excel va en inglés: el país también
        fac, env = c.datos_facturacion(), c.datos_envio()
    contacto = " · ".join(x for x in (c.persona_contacto, c.telefono) if x)
    # Billing address: una línea por fila, sin las vacías
    direccion = [x for x in ([fac["razon_social"]] + _lineas_lugar(fac)
                             + [f"VAT: {c.cif}" if c.cif else "", f"EORI: {c.eori}" if (ue and c.eori) else "",
                                f"Attn: {contacto}" if contacto else ""]) if x]
    recibe = " · ".join(x for x in (env.get("contacto"), env.get("telefono")) if x)
    entrega = [x for x in ([env["destinatario"]] + _lineas_lugar(env) + [f"Attn: {recibe}" if recibe else ""]) if x]
    t = totales(pedido, regimen, envio)
    return {
        "formato": formato, "regimen": regimen, "configurado": bool(c.formato_factura and c.regimen_iva),
        "envio_pactado": envio_pactado, "envio_cambiado": envio != envio_pactado,
        "numero": pedido.referencia_interna or f"#{pedido.id}", "fecha": hoy, "direccion": direccion,
        "entrega": entrega if env.get("direccion") else direccion, "ue": ue,
        "falta_direccion": not (c.direccion and c.cif),
        "referencia_cliente": pedido.referencia or "", "transportista": pedido.transportista or "",
        "uso": USO_EXCEL.get(pedido.uso, ""),
        "columnas": COLUMNAS[formato], "filas": filas_productos(pedido),
        "totales": t, "lista_totales": lista_totales(t),
        "avisos": avisos_kits(pedido) if formato != "A" else [],
    }


# --------------------------------------------------------------------------- #
# La factura en Excel
# --------------------------------------------------------------------------- #
# Columnas de cada formato: (clave, anchos de sus columnas de Excel). La primera de A y B («PRODUCT»)
# y la de C («Item») ocupan tres columnas, como en la plantilla. La última es la de los importes.
DISENO = {
    "A": [("producto", (3, 9, 8)), ("unidad", (6,)), ("referencia", (14,)), ("item", (34,)), ("cantidad", (8,)),
          ("hs", (30,))],
    "B": [("producto", (3, 9, 8)), ("unidad", (6,)), ("referencia", (14,)), ("item", (34,)), ("cantidad", (8,)),
          ("hs", (14,)), ("precio", (16,))],
    "C": [("item", (3, 16, 22)), ("referencia", (14,)), ("hs", (13,)), ("unitario", (12,)), ("cantidad", (8,)),
          ("precio", (16,))],
}
IMG = os.path.join(os.path.dirname(os.path.dirname(__file__)), "static", "img")


def _imagen(ws, archivo, celda, alto_px):
    """Pone una imagen de static/img con ese alto (si no se puede leer, no pasa nada: la factura sale igual)."""
    from openpyxl.drawing.image import Image
    ruta = os.path.join(IMG, archivo)
    try:
        img = Image(ruta)
    except Exception:
        from flask import current_app
        current_app.logger.warning("Factura en Excel: no se ha podido poner %s", ruta)
        return
    if alto_px:
        img.width, img.height = round(img.width * alto_px / img.height), alto_px
    ws.add_image(img, celda)


def excel_para_facturar(datos, empresa) -> bytes:
    """La factura en Excel (una hoja, en inglés, lista para pasar a PDF).
    empresa: los Datos de Yikon (datos_empresa.datos())."""
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    from app.excel import FORMATO_EUROS

    formato, columnas = datos["formato"], datos["columnas"]
    diseno = dict(DISENO[formato])
    col, n = {}, 1                                   # columna de Excel donde empieza cada dato
    for clave, anchos in DISENO[formato]:
        col[clave] = n
        n += len(anchos)
    ultima = n - 1
    col_valor = ultima                               # importes de los totales: la última columna
    L = get_column_letter

    libro = Workbook()
    ws = libro.active
    ws.title = "INVOICE"
    ws.sheet_view.showGridLines = False
    base, negrita = Font(size=10), Font(size=10, bold=True)
    cabecera_fondo = PatternFill("solid", fgColor="EFE3EE")
    linea = Side(style="thin", color="8C8C8C")
    borde = Border(top=linea, bottom=linea, left=linea, right=linea)
    for clave, anchos in DISENO[formato]:
        for k, ancho in enumerate(anchos):
            ws.column_dimensions[L(col[clave] + k)].width = ancho

    def escribir(fila, columna, valor, fuente=None, **alineado):
        celda = ws.cell(row=fila, column=columna, value=valor)
        celda.font = fuente or base
        if alineado:
            celda.alignment = Alignment(**alineado)
        return celda

    # 1. Logo (el de la plantilla, ya reducido en el archivo: Excel no lo escala y se ve nítido) y «INVOICE»
    for f in range(1, 6):
        ws.row_dimensions[f].height = 14
    ws.row_dimensions[3].height = 26
    _imagen(ws, "factura-logo.png", "B1", None)
    escribir(3, ultima, "INVOICE", Font(size=18, bold=True, color="A61C92"), horizontal="right", vertical="center")

    # 2. Seller / Billing Address / Delivery address, uno al lado del otro
    vendedor = [empresa.get("razon_social") or empresa.get("nombre")] + (empresa.get("direccion") or "").splitlines()
    vendedor += [x for x in (empresa.get("cp_ciudad"), (empresa.get("pais") or "").upper()) if x]
    if empresa.get("cif"):
        vendedor.append(f"CIF: {empresa['cif']}")
    if datos["ue"] and empresa.get("eori"):
        vendedor.append(f"EORI: {empresa['eori']}")
    vendedor = [x for x in vendedor if x]
    if formato == "C":
        inicios = [2, col["referencia"], col["cantidad"]]
    else:
        inicios = [2, col["item"], col["hs"]]
    bloques = [("Seller.", vendedor), ("Billing Address.", datos["direccion"]), ("Delivery address", datos["entrega"])]
    # Cada bloque tiene el ancho hasta el siguiente (el último, hasta el borde de la hoja): las líneas
    # más largas se parten en dos para que no se corten ni se salgan de la página
    import textwrap
    anchos_col = {}
    for clave, anchos in DISENO[formato]:
        for k, a in enumerate(anchos):
            anchos_col[col[clave] + k] = a
    limites = inicios[1:] + [ultima + 1]
    partidos = []
    for (titulo, lineas), c0, fin in zip(bloques, inicios, limites):
        caben = max(12, int(sum(anchos_col.get(c, 9) for c in range(c0, fin)) * 1.1) - 1)
        partidos.append((titulo, [t for linea in lineas for t in (textwrap.wrap(linea, caben) or [linea])]))
    bloques = partidos
    for (titulo, lineas), c0 in zip(bloques, inicios):
        escribir(6, c0, titulo, negrita)
        for i, texto in enumerate(lineas, start=7):
            escribir(i, c0, texto)
    fila = 7 + max(len(b[1]) for b in bloques)
    # Fecha, nº de factura y su nº de pedido, debajo de la dirección de facturación
    for texto in (f"Invoice Date: {datos['fecha']:%d/%m/%Y}", f"Invoice Number: {datos['numero']}",
                  f"Your PO: {datos['referencia_cliente']}" if datos["referencia_cliente"] else ""):
        if texto:
            fila += 1
            escribir(fila, inicios[1], texto, negrita)
    fila += 2

    # 3. Los productos
    for clave in columnas:
        nombre = "Total Price" if (clave == "precio" and formato == "C") else COLUMNAS_EXCEL[clave]
        celda = escribir(fila, col[clave], nombre, negrita, horizontal="center", vertical="center", wrap_text=True)
        celda.fill = cabecera_fondo
        ancho = len(diseno[clave])
        for k in range(ancho):
            ws.cell(row=fila, column=col[clave] + k).border = borde
            ws.cell(row=fila, column=col[clave] + k).fill = cabecera_fondo
        if ancho > 1:
            ws.merge_cells(start_row=fila, start_column=col[clave], end_row=fila, end_column=col[clave] + ancho - 1)
    ws.row_dimensions[fila].height = 22
    primera = fila + 1
    agrupado = formato in ("A", "B")
    for f in datos["filas"]:
        fila += 1
        for clave in columnas:
            ancho = len(diseno[clave])
            for k in range(ancho):
                ws.cell(row=fila, column=col[clave] + k).border = borde
            if agrupado and clave in ("producto", "unidad") and not f["primera"]:
                continue                                         # el kit y cuántos: solo en su primera fila
            valor = f[clave]
            if clave in ("unitario", "precio"):
                valor = float(valor)
            elif clave == "hs" and str(valor).isdigit():
                valor = int(valor)                               # (como en la plantilla: número)
            celda = escribir(fila, col[clave], valor, negrita if (clave == "producto") else None,
                             vertical="center", wrap_text=clave in ("producto", "item"),
                             horizontal="center" if clave in ("unidad", "cantidad", "hs") else None)
            if clave in ("unitario", "precio"):
                celda.number_format = FORMATO_EUROS
            alto_grupo = f["grupo"] if (agrupado and clave in ("producto", "unidad")) else 1
            if ancho > 1 or alto_grupo > 1:
                ws.merge_cells(start_row=fila, start_column=col[clave], end_row=fila + alto_grupo - 1,
                               end_column=col[clave] + ancho - 1)
    ultima_fila = fila

    # 4. Totales: nombre a la izquierda de la columna de importes; fórmulas donde se puede
    letra = L(col_valor)
    refs, t = {}, datos["totales"]
    for clave, valor in datos["lista_totales"]:
        fila += 1
        etiqueta = "Total before discount" if (clave == "bruto" and formato == "C") else TOTALES_EXCEL[clave]
        escribir(fila, col_valor - 1, etiqueta + ":", negrita, horizontal="right")
        celda = ws.cell(row=fila, column=col_valor)
        celda.font = negrita
        refs[clave] = f"{letra}{fila}"
        neto = refs.get("neto", refs.get("bruto"))
        if clave == "bruto":
            celda.value = (f"=SUM({letra}{primera}:{letra}{ultima_fila})" if "precio" in columnas and datos["filas"]
                           else float(valor))
        elif clave == "descuento":
            celda.value, celda.number_format = float(valor), '0.##" %"'
        elif clave == "neto":
            celda.value = f"={refs['bruto']}*(1-{refs['descuento']}/100)"
        elif clave == "envio":
            celda.value = float(valor)
        elif clave == "iva_pct":
            # Varios tipos de IVA: no hay un % que aplicar; el importe va calculado
            celda.value = float(valor) if valor is not None else "mixed"
            celda.number_format = '0.##" %"'
        elif clave == "iva":
            if t["iva_pct"] is not None:
                base_iva = neto + (f"+{refs['envio']}" if "envio" in refs else "")
                celda.value = f"=ROUND(({base_iva})*{refs['iva_pct']}/100,2)"
            else:
                celda.value = float(valor)
        elif clave == "total":
            partes = [neto] + [refs[k] for k in ("envio", "iva") if k in refs]
            celda.value = "=" + "+".join(partes)
            celda.font = Font(size=11, bold=True)
            ws.cell(row=fila, column=col_valor - 1).font = Font(size=11, bold=True)
        if clave not in ("descuento", "iva_pct"):
            celda.number_format = FORMATO_EUROS
        celda.alignment = Alignment(horizontal="right")
    if datos["uso"]:
        fila += 2
        escribir(fila, 2, f"USE: {datos['uso']}", negrita)

    # 5. Transferencia, entrega, Incoterm y transportista (como la plantilla de Yikon)
    fila += 2
    escribir(fila, 1, "Wire Transfer", negrita)
    fila += 1
    escribir(fila, 2, "INSTRUCTIONS FOR USD/EUR/RMB PAYMENTS:")
    _imagen(ws, "factura-transferencia.png", f"B{fila + 1}", 60)
    fila += 5
    nombre_empresa = (empresa.get("razon_social") or empresa.get("nombre") or "").upper()
    instrucciones = [
        ("PLEASE REQUEST THE REMITTING BANK TO CLEARLY AND ACCURATELY STATE THE FOLLOWING INFORMATION "
         "WHEN PERFORMING THE WIRE TRANSFER:", None),
        ("1. YOUR (SENDER’S) INFORMATION (COMPANY NAME, TELEPHONE, ETC.)", None),
        ("2. WIRE TRANSFER AMOUNT (PLEASE CLEARLY STATE THAT “ALL ADDITIONAL FEES ASSOCIATED WITH THIS TRANSFER "
         "WILL BE CHARGED TO THE REMITTING BANK.”)", None),
        (f"3. {nombre_empresa} BENEFICIARY BANK:", None),
        (f"BENEFICIARY NAME: {nombre_empresa}", negrita),
        (f"BNF’S ACCOUNT NO.: {empresa.get('iban') or ''}", negrita),
        (f"NAME OF THE BANK: {empresa.get('banco') or ''}", negrita),
        (f"SWIFT CODE: {empresa.get('swift') or ''}", negrita),
        (f"4. REFERENCE CONTRACT NUMBER: {datos['referencia_cliente'] or datos['numero']}", None),
        (f"5. DELIVERY: {empresa.get('entrega') or ''}", None),
        (f"6. Incoterm: {empresa.get('incoterm') or ''}", None),
        (f"7. Shipped by: {datos['transportista']}", None),
    ]
    for texto, fuente in instrucciones:
        celda = escribir(fila, 2, texto, fuente, wrap_text=True, vertical="top")
        ws.merge_cells(start_row=fila, start_column=2, end_row=fila, end_column=ultima)
        if len(texto) > 95:
            ws.row_dimensions[fila].height = 27
        fila += 1

    # 6. Firmas: el cliente a la izquierda; Yikon (con su sello) a la derecha
    fila += 1
    escribir(fila, 2, "Confirmed & Accepted By:", negrita)
    for i, texto in enumerate(("For and on behalf of", nombre_empresa)):
        escribir(fila + i, inicios[2], texto, negrita, wrap_text=True)
        ws.merge_cells(start_row=fila + i, start_column=inicios[2], end_row=fila + i, end_column=ultima)
    ws.row_dimensions[fila + 2].height = 62
    _imagen(ws, "sello-yikon.png", f"{L(inicios[2])}{fila + 2}", 78)
    escribir(fila + 3, inicios[2], "Authorized Signature", negrita)
    for c in range(inicios[2], ultima + 1):
        ws.cell(row=fila + 3, column=c).border = Border(top=linea)
    fila += 5
    escribir(fila, 1, "THANK YOU FOR YOUR BUSINESS!", negrita, horizontal="center")
    ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=ultima)
    contacto = "   ".join(x for x in (f"CUSTOMER SERVICE: {empresa['email']}" if empresa.get("email") else "",
                                      f"WEBSITE: {empresa['web']}" if empresa.get("web") else "") if x)
    if contacto:
        fila += 1
        escribir(fila, 1, contacto, horizontal="center")
        ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=ultima)

    # Para imprimirla o pasarla a PDF: A4 vertical, todo a lo ancho de una página
    ws.print_area = f"A1:{L(ultima)}{fila}"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    # (con pocas líneas, entera en una página; con muchas, a lo ancho y las páginas que hagan falta)
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, (1 if len(datos["filas"]) <= 15 else 0)
    ws.page_margins.left = ws.page_margins.right = 0.4
    # Centrada en la página: al encajarla entera en una hoja se encoge y, sin esto, sobraba hueco a la derecha
    ws.print_options.horizontalCentered = True

    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()
