"""
Exportar tablas a Excel (.xlsx) con openpyxl (Python puro: sin DLL, funciona en tu Windows).

Una sola función, `generar_excel`, que sirve para cualquier listado: recibe las columnas
y las filas ya preparadas. Así pedidos, facturas o lo que venga después se exportan igual:
cabecera con el color corporativo, filtros en la cabecera, primera fila fija al hacer
scroll, fechas como fechas y euros como números (se pueden sumar en Excel).
"""
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

MARCA = "A61C92"
FORMATO_EUROS = '#,##0.00 "€"'
FORMATO_FECHA = "DD/MM/YYYY"


class Columna:
    """titulo: cabecera · ancho: en caracteres · tipo: "texto", "fecha", "euros" o "numero"."""

    def __init__(self, titulo, ancho=15, tipo="texto"):
        self.titulo, self.ancho, self.tipo = str(titulo), ancho, tipo


def generar_excel(hoja, columnas, filas, subtitulo=None) -> bytes:
    """
    hoja: nombre de la pestaña · columnas: [Columna, ...] · filas: [[valor, ...], ...]
    subtitulo: texto opcional encima de la tabla (p. ej. los filtros aplicados).
    Devuelve el archivo .xlsx en bytes.
    """
    libro = Workbook()
    ws = libro.active
    ws.title = str(hoja)[:31]                    # Excel no admite nombres de pestaña más largos

    fila_cabecera = 1
    if subtitulo:
        ws.cell(row=1, column=1, value=str(subtitulo)).font = Font(italic=True, color="666666")
        fila_cabecera = 3

    for i, col in enumerate(columnas, start=1):
        celda = ws.cell(row=fila_cabecera, column=i, value=col.titulo)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor=MARCA)
        celda.alignment = Alignment(vertical="center", horizontal="right" if col.tipo in ("euros", "numero") else "left")
        ws.column_dimensions[get_column_letter(i)].width = col.ancho

    for n, fila in enumerate(filas, start=fila_cabecera + 1):
        for i, (col, valor) in enumerate(zip(columnas, fila), start=1):
            if isinstance(valor, Decimal):
                valor = float(valor)                 # Excel guarda números, no Decimal
            elif isinstance(valor, datetime):
                valor = valor.replace(tzinfo=None)   # Excel no entiende zonas horarias
            celda = ws.cell(row=n, column=i, value=valor)
            if isinstance(valor, str) and valor.startswith("="):
                # Un texto que empieza por "=" (p. ej. una referencia escrita por un cliente)
                # Excel lo ejecutaría como fórmula al abrir el archivo: lo forzamos a texto.
                celda.data_type = "s"
            if col.tipo == "euros":
                celda.number_format = FORMATO_EUROS
            elif col.tipo == "fecha" and isinstance(valor, (date, datetime)):
                celda.number_format = FORMATO_FECHA
            if isinstance(valor, str) and "\n" in valor:
                celda.alignment = Alignment(wrap_text=True, vertical="top")
            else:
                celda.alignment = Alignment(vertical="top")

    ultima = fila_cabecera + max(len(filas), 1)
    ws.auto_filter.ref = f"A{fila_cabecera}:{get_column_letter(len(columnas))}{ultima}"
    ws.freeze_panes = ws.cell(row=fila_cabecera + 1, column=1)   # la cabecera queda fija

    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()


MIMETYPE_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
