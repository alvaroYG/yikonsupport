"""¿Dónde está este lote?: a qué clientes y pedidos fue cada lote (y qué queda en la oficina)."""
from datetime import date, timedelta

from app.extensions import db
from app.models import Pedido
from tests.conftest import llevar_a_validado, login, pedido_cliente

CAD = (date.today() + timedelta(days=200)).isoformat()


def _enviar_pedido(client, cliente, fas, lote_id=None):
    """El cliente pide 2 × P1; se prepara, valida y envía (opcionalmente eligiendo el lote)."""
    login(client, cliente); pedido_cliente(client, cantidad=2); client.post("/auth/logout")
    p = db.session.scalar(db.select(Pedido).order_by(Pedido.id.desc()))
    login(client, "admin"); client.post(f"/pedidos/{p.id}/preparacion"); client.post("/auth/logout")
    llevar_a_validado(client, p.id, fas)
    login(client, "admin")
    datos = {"estado": "enviado", "confirmar_stock": "1", "transportista": "SEUR", "numero_seguimiento": "T1"}
    client.post(f"/pedidos/{p.id}/estado", data=datos)
    client.post("/auth/logout")
    return p


def _escenario(client):
    login(client, "admin")
    client.post("/inventario/movimiento", data={"articulo": "producto-1", "tipo": "entrada", "cantidad": 10,
                                                "lote": "YK-RETIRADO", "caducidad": CAD})
    client.post("/auth/logout")
    pa = _enviar_pedido(client, "cliA", "loga")
    pb = _enviar_pedido(client, "cliB", "logb")
    return pa, pb


def test_admin_ve_a_quien_fue_el_lote_y_lo_que_queda(app, client):
    pa, pb = _escenario(client)
    login(client, "admin")
    html = client.get("/inventario/rastrear-lote?q=yk-retir").get_data(as_text=True)     # parte del código, sin mayúsculas
    assert "YK-RETIRADO" in html and "Clínica A" in html and "Clínica B" in html
    assert pa.referencia_interna in html and pb.referencia_interna in html
    assert "En la oficina" in html and "2 clientes" in html
    from app.inventario.rastreo import buscar_lotes, rastro
    with client:                                                    # (current_user = admin, dentro de una petición)
        client.get("/")
        r = rastro(buscar_lotes("YK-RETIRADO")[0], con_oficina=True)
        assert (r["entradas"], r["enviadas"], r["en_oficina"]) == (10, 4, 6)
    assert "No hay ningún lote" in client.get("/inventario/rastrear-lote?q=NOEXISTE").get_data(as_text=True)


def test_un_envio_anulado_no_cuenta(app, client):
    pa, pb = _escenario(client)
    login(client, "admin")
    client.post(f"/pedidos/{pb.id}/estado", data={"estado": "validado"})          # se anula el envío de B
    html = client.get("/inventario/rastrear-lote?q=YK-RETIRADO").get_data(as_text=True)
    # (se mira la referencia del pedido: el nombre de la clínica sale en el desplegable de "nuevo pedido")
    assert pa.referencia_interna in html and pb.referencia_interna not in html and "1 cliente" in html


def test_el_fas_solo_ve_los_envios_a_sus_clientes(app, client):
    pa, pb = _escenario(client)
    login(client, "loga")
    html = client.get("/inventario/rastrear-lote?q=YK-RETIRADO").get_data(as_text=True)
    assert "Clínica A" in html and "Clínica B" not in html and pb.referencia_interna not in html
    assert "En la oficina" not in html                             # no ve el inventario
    client.post("/auth/logout")
    # un lote que nunca fue a sus clientes: para él no existe
    login(client, "admin")
    client.post("/inventario/movimiento", data={"articulo": "producto-1", "tipo": "entrada", "cantidad": 3,
                                                "lote": "SOLO-OFICINA", "caducidad": CAD})
    client.post("/auth/logout")
    login(client, "loga")
    assert "No hay ningún lote" in client.get("/inventario/rastrear-lote?q=SOLO-OFICINA").get_data(as_text=True)


def test_excel_y_permisos(app, client):
    from io import BytesIO
    from openpyxl import load_workbook
    _escenario(client)
    login(client, "ofi")                                           # Logística también rastrea
    r = client.get("/inventario/rastrear-lote.xlsx?q=YK-RETIRADO")
    assert r.status_code == 200
    filas = [[c.value for c in f] for f in load_workbook(BytesIO(r.data)).active.iter_rows()]
    clientes = {f[3] for f in filas if f[0] == "YK-RETIRADO"}
    assert clientes == {"Clínica A", "Clínica B"}
    client.post("/auth/logout")
    for usuario in ("cliA", "sop", "chrA"):                        # el cliente nunca ve lotes
        login(client, usuario)
        assert client.get("/inventario/rastrear-lote?q=YK").status_code == 403, usuario
        assert client.get("/inventario/rastrear-lote.xlsx?q=YK").status_code == 403, usuario
        client.post("/auth/logout")


def test_al_fas_no_se_le_escapa_su_lote_entre_muchos(app, client):
    """Revisión: con cientos de lotes parecidos, el FAS sigue encontrando el que llegó a su cliente
    (se parte de sus envíos, no de "los 200 primeros")."""
    from app.models import Lote
    _escenario(client)                                     # YK-RETIRADO enviado a las clínicas A y B
    db.session.add_all([Lote(producto_id=1, codigo=f"AAA-{i:04d}-YK") for i in range(260)])
    db.session.commit()
    login(client, "loga")
    html = client.get("/inventario/rastrear-lote?q=YK").get_data(as_text=True)
    assert "YK-RETIRADO" in html and "AAA-0001-YK" not in html


def test_el_rastreo_hace_pocas_consultas(app, client):
    """Revisión: el buscador pide sugerencias en cada pulsación; no puede lanzar cientos de consultas."""
    from sqlalchemy import event
    from app.models import Lote
    _escenario(client)
    db.session.add_all([Lote(producto_id=1, codigo=f"YK-M-{i:04d}") for i in range(150)])
    db.session.commit()
    login(client, "admin")
    consultas = []

    def contar(*_args):
        consultas.append(1)

    event.listen(db.engine, "before_cursor_execute", contar)
    try:
        assert client.get("/buscar/sugerencias?q=YK").status_code == 200
    finally:
        event.remove(db.engine, "before_cursor_execute", contar)
    assert len(consultas) < 60, len(consultas)
