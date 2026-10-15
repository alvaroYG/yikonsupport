"""Fallos encontrados en la revisión del 27/09 (lotes, envíos, Excel, invitaciones...)."""
import re
from datetime import date, timedelta
from io import BytesIO

from openpyxl import load_workbook

from app.extensions import db
from app.models import Lote, MovimientoInventario, Pedido, Producto, User
from tests.conftest import doc_pedido, confirmar_correos, contenido, llevar_a_validado, login, pedido_cliente
from tests.test_inventario import _entrada, _estado, _pedido_confirmado


def test_excel_no_ejecuta_formulas_escritas_por_clientes(app, client):
    login(client, "cliA")
    pedido_cliente(client, referencia='=HYPERLINK("http://malo.example","Clic")')
    client.post("/auth/logout"); login(client, "admin")
    ws = load_workbook(BytesIO(client.get("/pedidos/exportar.xlsx").data)).active
    celda = next(c for fila in ws.iter_rows() for c in fila if str(c.value).startswith("=HYPERLINK"))
    assert celda.data_type == "s"                           # texto, no fórmula


def test_la_invitacion_no_sirve_si_ya_entro(app, client):
    correos = app.extensions.setdefault("correos_enviados", []); correos.clear()
    login(client, "admin")
    client.post("/admin/usuarios/nuevo", data={"nombre": "Eva", "email": "eva@ejemplo.com", "rol": "admin",
                                               "cliente_id": 0, "activo": "y", "invitar": "y",
                                               "password": "contraseña-del-admin"})
    assert confirmar_correos() == 1
    enlace = re.search(r"/auth/invitacion/[^\s]+", correos[0].get_content()).group(0)
    client.post("/auth/logout")
    client.post("/auth/login", data={"email": "eva@ejemplo.com", "password": "contraseña-del-admin"})
    client.post("/auth/logout")
    r = client.post(enlace, data={"nueva": "otra-contraseña-1", "confirmar": "otra-contraseña-1"},
                    follow_redirects=True)
    assert "no es válida o ha caducado" in r.get_data(as_text=True)
    assert db.session.scalar(db.select(User).filter_by(email="eva@ejemplo.com")).check_password("contraseña-del-admin")


def test_abrir_un_enlace_de_invitacion_falso_no_cierra_la_sesion(app, client):
    login(client, "admin")
    client.get("/auth/invitacion/basura")
    assert client.get("/admin/usuarios").status_code == 200


def test_enviar_con_enlace_de_seguimiento_invalido_explica_el_error(app, client):
    p = _pedido_confirmado(client, 1)
    llevar_a_validado(client, p.id)
    login(client, "admin"); _entrada(client, 5)
    r = client.post(f"/pedidos/{p.id}/estado", data={"estado": "enviado", "transportista": "SEUR", "numero_seguimiento": "TRK1", "confirmar_stock": "1",
                                                     "enlace_seguimiento": "www.seur.com/x"}, follow_redirects=True)
    assert r.status_code == 200 and "https://" in r.get_data(as_text=True)
    assert db.session.get(Pedido, p.id).estado == "validado"


def test_stock_negativo_heredado_no_permite_enviar_de_mas(app, client):
    # Situación que dejaba la versión anterior: un ajuste que llevó el stock a −2 ("SIN LOTE")
    prod = db.session.get(Producto, 1)
    viejo = Lote(producto=prod, codigo="SIN LOTE", caducidad=None)
    db.session.add_all([viejo, MovimientoInventario(producto=prod, lote=viejo, cantidad=-2, tipo="ajuste")])
    db.session.commit()
    login(client, "admin"); _entrada(client, 3, lote="A", caducidad="2099-01-01"); client.post("/auth/logout")
    p = _pedido_confirmado(client, 3)
    llevar_a_validado(client, p.id)
    login(client, "admin")
    _estado(client, p, "enviado")
    assert db.session.get(Pedido, p.id).estado == "validado"          # solo hay 1 de verdad
    assert db.session.get(Producto, 1).stock == 1


def test_referencia_de_solo_espacios_en_pedido_interno(app, client):
    login(client, "loga")
    for _i in range(2):
        r = client.post("/pedidos/nuevo", data={"uso": "ruo", "documento": doc_pedido(), "cliente_id": 1, "fecha": "2026-09-01", "referencia": "   ",
                                                "contenido": contenido(sueltos={1: 1})})
        assert r.status_code == 302
    assert [p.referencia for p in db.session.scalars(db.select(Pedido))] == [None, None]


def test_no_se_mete_mercancia_nueva_en_el_lote_sin_caducidad(app, client):
    prod = db.session.get(Producto, 1)
    db.session.add(Lote(producto=prod, codigo="SIN LOTE", caducidad=None)); db.session.commit()
    login(client, "admin")
    _entrada(client, 4, lote="SIN LOTE", caducidad=(date.today() + timedelta(days=2)).isoformat())
    assert db.session.get(Producto, 1).stock == 0


def test_repetir_un_pedido_de_mas_de_20_productos(app, client):
    prods = [Producto(referencia=f"R{i}", nombre=f"Prod {i}", precio=1) for i in range(22)]
    db.session.add_all(prods); db.session.commit()
    login(client, "loga")
    datos = {"cliente_id": 1, "fecha": "2026-09-01",
             "contenido": contenido(sueltos={p.id: 1 for p in prods[:20]}), "documento": doc_pedido()}
    client.post("/pedidos/nuevo", data=datos)
    pedido = db.session.scalar(db.select(Pedido))
    for p in prods[20:]:
        client.post(f"/pedidos/{pedido.id}/lineas", data={"producto_id": p.id, "cantidad": 1, "descuento_pct": 0})
    assert len(pedido.lineas) == 22
    assert client.get(f"/pedidos/nuevo?repetir={pedido.id}").status_code == 200


def test_estadisticas_no_cuentan_pedidos_con_fecha_futura(app, client):
    login(client, "cliA")
    pedido_cliente(client)
    client.post("/pedidos/nuevo", data={"uso": "ruo", "documento": doc_pedido(), "contenido": contenido([(1, 1, {1: 1})]),
                                        "fecha": (date.today() + timedelta(days=60)).isoformat(),
                                        "referencia": "FUTURO"})
    client.post("/auth/logout"); login(client, "admin")
    html = client.get("/estadisticas").get_data(as_text=True)
    assert re.search(r'etiqueta">Pedidos</div><div class="valor">1<', html)


def test_la_ficha_de_un_pedido_enviado_no_repite_ids(app, client):
    p = _pedido_confirmado(client, 1)
    llevar_a_validado(client, p.id)
    login(client, "admin"); _entrada(client, 5); _estado(client, p, "enviado")
    html = client.get(f"/pedidos/{p.id}").get_data(as_text=True)
    ids = re.findall(r'\sid="([^"]+)"', html)
    assert len(ids) == len(set(ids)), [i for i in ids if ids.count(i) > 1]


def test_reiniciar_datos_deja_solo_los_admin(app, client):
    """flask reiniciar-datos: borra todo menos los usuarios admin."""
    from app.models import Cliente, Kit, MovimientoInventario, Producto, User
    login(client, "cliA"); pedido_cliente(client); client.post("/auth/logout")
    login(client, "admin"); _entrada(client, 3); client.post("/auth/logout")
    admins_antes = {u.email for u in db.session.scalars(db.select(User).filter_by(rol="admin"))}
    r = app.test_cli_runner().invoke(args=["reiniciar-datos", "--si"])
    assert "Hecho" in r.output, r.output
    db.session.expire_all()
    assert {u.email for u in db.session.scalars(db.select(User))} == admins_antes
    for modelo in (Cliente, Pedido, Producto, Kit, MovimientoInventario):
        assert db.session.scalar(db.select(db.func.count()).select_from(modelo)) == 0
    login(client, "admin")                                   # el admin sigue pudiendo entrar
    assert client.get("/").status_code == 200


def test_reiniciar_datos_pregunta_antes(app, client):
    from app.models import Cliente
    r = app.test_cli_runner().invoke(args=["reiniciar-datos"], input="n\n")
    assert "Cancelado" in r.output
    assert db.session.scalar(db.select(db.func.count(Cliente.id))) > 0
