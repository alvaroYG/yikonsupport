"""flask importar-datos: carga de productos, kits y empresas desde la plantilla Excel."""
import os
from decimal import Decimal

from openpyxl import Workbook, load_workbook

from app.extensions import db
from app.importar import COLUMNAS, _sin_tildes
from app.models import Cliente, Kit, Producto, RegistroAuditoria

CAB = {
    "Productos": ["Referencia *", "Nombre *", "Marca", "Precio € (sin IVA) *", "IVA % *", "Stock mínimo",
                  "Controlar stock (Sí/No)", "Activo (Sí/No)", "Notas (no se carga)"],
    "Kits": ["Referencia kit *", "Nombre *", "Descripción", "Precio € (sin IVA) *", "IVA % *",
             "Tipo (Montado/Cerrado) *", "Activo (Sí/No)", "Notas (no se carga)"],
    "Componentes": ["Referencia kit *", "Referencia producto *", "Cantidad por kit *", "Notas (no se carga)"],
    "Empresas": ["Nombre *", "Razón social", "CIF", "Prefijo de pedidos", "Email", "Teléfono", "Dirección",
                 "Ciudad", "Persona de contacto", "IVA (España / UE / Exportación)", "Formato de factura (A/B/C)",
                 "Gastos de envío €", "Ve precios (Sí/No)", "Revisar facturas (Sí/No)", "Soporte de cliente (email)", "Notas"],
}


def _excel(tmp_path, **hojas):
    libro = Workbook()
    libro.remove(libro.active)
    for nombre, filas in hojas.items():
        ws = libro.create_sheet(nombre)
        ws.append(CAB[nombre])
        for fila in filas:
            ws.append(fila)
    ruta = os.path.join(tmp_path, "carga.xlsx")
    libro.save(ruta)
    return ruta


def _importar(app, ruta, guardar=False):
    return app.test_cli_runner().invoke(args=["importar-datos", ruta] + (["--guardar"] if guardar else []))


BUENO = dict(
    Productos=[["XK-043", "Sample Preservation Buffer", "Yikon", "12,50", 10, 3, "Sí", "Sí", "nota que no se carga"],
               ["YK001-001", "NGS Universal Index Kit", None, None, None, None, "No", None, None]],
    Kits=[["PGTA-24", "PGT-A · 24 rxn", "Para 24 muestras", 1500, "21 %", "Montado", "Sí", None],
          ["PGTA-96", "PGT-A · 96 rxn", None, None, None, "Cerrado", "No", None]],
    Componentes=[["PGTA-24", "XK-043", 2, None], ["PGTA-24", "yk001-001", 1, None],     # (da igual mayúsculas)
                 ["PGTA-96", "XK-043", 8, None], ["PGTA-96", "P1", 1, None]],            # P1 ya estaba en la plataforma
    Empresas=[["Clínica Nueva", "Clínica Nueva S.L.", "b-123 456 78", None, "Info@Nueva.es", None, "C/ Mayor 1",
               "Valencia", "Marta", "UE", "b", "15", "Sí", None, "loga@ejemplo.com", None],
              ["Laboratorio Sur", None, None, "LSUR", None, None, None, None, None, None, None, None, None, "Sí", None, None]],
)


def test_la_prueba_no_guarda_nada(app, tmp_path):
    r = _importar(app, _excel(tmp_path, **BUENO))
    assert r.exit_code == 0, r.output
    assert "Productos: 2 nuevos (1 sin precio)" in r.output and "Kits: 2 nuevos (1 sin precio, 4 componentes)" in r.output
    assert "Empresas: 2 nuevas" in r.output and "PRUEBA" in r.output
    assert db.session.scalar(db.select(db.func.count(Producto.id))) == 1            # solo el P1 del escenario
    assert db.session.scalar(db.select(Kit).filter_by(referencia="PGTA-24")) is None
    assert db.session.scalar(db.select(db.func.count(Cliente.id))) == 2


def test_guardar_lo_crea_todo_con_sus_datos(app, tmp_path):
    r = _importar(app, _excel(tmp_path, **BUENO), guardar=True)
    assert r.exit_code == 0 and "Guardado." in r.output, r.output
    xk = db.session.scalar(db.select(Producto).filter_by(referencia="XK-043"))
    assert (xk.nombre, xk.categoria, xk.precio, xk.iva_pct, xk.stock_minimo) == \
        ("Sample Preservation Buffer", "Yikon", Decimal("12.50"), Decimal("10.00"), 3)
    assert xk.controlar_stock and xk.activo
    yk = db.session.scalar(db.select(Producto).filter_by(referencia="YK001-001"))
    assert yk.precio == 0 and yk.iva_pct == Decimal("21.00") and not yk.controlar_stock   # vacíos: 0 € y 21 %
    k24 = db.session.scalar(db.select(Kit).filter_by(referencia="PGTA-24"))
    assert (k24.precio, k24.iva_pct, k24.cerrado, k24.activo, k24.descripcion) == \
        (Decimal("1500.00"), Decimal("21.00"), False, True, "Para 24 muestras")
    assert [(c.producto.referencia, c.cantidad) for c in k24.componentes] == [("XK-043", 2), ("YK001-001", 1)]
    k96 = db.session.scalar(db.select(Kit).filter_by(referencia="PGTA-96"))
    assert k96.cerrado and not k96.activo and [c.producto.referencia for c in k96.componentes] == ["XK-043", "P1"]
    nueva = db.session.scalar(db.select(Cliente).filter_by(nombre="Clínica Nueva"))
    assert (nueva.cif, nueva.prefijo, nueva.email, nueva.regimen_iva, nueva.formato_factura) == \
        ("B12345678", "Clinica", "info@nueva.es", "ue", "B")                      # prefijo sacado del nombre
    assert nueva.gastos_envio == Decimal("15.00") and nueva.ver_precios and not nueva.revisar_facturas
    assert [u.email for u in nueva.logisticos] == ["loga@ejemplo.com"]               # su Soporte de cliente
    sur = db.session.scalar(db.select(Cliente).filter_by(nombre="Laboratorio Sur"))
    assert sur.prefijo == "LSUR" and sur.revisar_facturas and not sur.ver_precios and sur.logisticos == []
    # queda en la auditoría
    detalles = db.session.scalars(db.select(RegistroAuditoria.detalle).filter_by(accion="crear")).all()
    assert sum("importado de carga.xlsx" in d for d in detalles) == 6
    assert any(d.startswith("PGTA-24 · 2 productos") for d in detalles)


def test_lo_que_ya_existe_no_se_toca(app, tmp_path):
    ruta = _excel(tmp_path,
                  Productos=[["p1", "Otro nombre", None, 999, 21, None, None, None, None],
                             ["NUEVO", "Producto nuevo", None, 5, 21, None, None, None, None]],
                  Kits=[["K1", "Otro kit", None, 1, 21, "Cerrado", None, None]],
                  Componentes=[["K1", "NUEVO", 3, None]],
                  Empresas=[["clínica a", None, None, "ZZZ", None, None, None, None, None, None, None, None, None, None, None, None],
                            ["Otra", None, "B11111111", "OTRA", None, None, None, None, None, None, None, None, None, None, None, None]])
    r = _importar(app, ruta, guardar=True)
    assert r.exit_code == 0, r.output
    assert "p1 ya existe: no se toca" in r.output and "K1 ya existe" in r.output
    assert "ya existe (como «Clínica A»)" in r.output                             # por nombre y, la otra, por CIF
    assert "Productos: 1 nuevos" in r.output and "Empresas: 0 nuevas, 2 ya existían" in r.output
    p1 = db.session.scalar(db.select(Producto).filter_by(referencia="P1"))
    assert p1.nombre == "Kit 1" and p1.precio == 100
    k1 = db.session.scalar(db.select(Kit).filter_by(referencia="K1"))
    assert k1.nombre == "Kit uno" and [c.producto.referencia for c in k1.componentes] == ["P1"]
    assert db.session.scalar(db.select(Cliente).filter_by(prefijo="ZZZ")) is None


def test_con_un_error_no_se_guarda_nada(app, tmp_path):
    ruta = _excel(tmp_path,
                  Productos=[["A1", "Bien", None, 1, 21, None, None, None, None],
                             ["A1", "Repetido", None, 1, 21, None, None, None, None],
                             ["A2", None, None, 1, 21, None, None, None, None],
                             ["A3", "IVA raro", None, 1, 7, None, None, None, None],
                             ["A4", "Precio raro", None, "caro", 21, None, None, None, None]],
                  Kits=[["KA", "Kit A", None, 1, 21, "Montado", None, None],
                        ["KB", "Kit sin productos", None, 1, 21, "Montado", None, None],
                        ["KC", "Sin tipo", None, 1, 21, None, None, None]],
                  Componentes=[["KA", "NO-EXISTE", 1, None], ["KA", "A1", 0, None], ["KX", "A1", 1, None]],
                  Empresas=[["Clínica Uno", None, None, None, None, None, None, None, None, None, None, None, None, None, None, None],
                            ["Clínica Dos", None, None, None, None, None, None, None, None, None, None, None, None, None, None, None],
                            ["Tercera", None, None, "A B", None, None, None, None, None, "Marte", "D", None, "Quizá",
                             None, "clia@ejemplo.com", None]])
    r = _importar(app, ruta, guardar=True)
    assert r.exit_code == 1 and "NO se ha guardado nada" in r.output
    for texto in ["Productos, fila 3: la referencia A1 está repetida en la hoja (fila 2)",
                  "Productos, fila 4: faltan la referencia o el nombre", "IVA 7: tiene que ser 21, 10, 4 o 0",
                  "«Precio» no es un número: caro", "Kits, fila 4: «Tipo» tiene que ser Montado o Cerrado",
                  "el producto NO-EXISTE no está", "«Cantidad» tiene que ser 1 o más",
                  "el kit KX no está en la hoja «Kits»", "el kit KB no tiene ningún producto",
                  "Empresas, fila 3: el prefijo «Clinica» ya lo usa otra empresa (sale del nombre",
                  "prefijo «A B»: de 2 a 10 letras"]:
        assert texto in r.output, texto
    assert db.session.scalar(db.select(Producto).filter_by(referencia="A1")) is None
    assert db.session.scalar(db.select(Kit).filter_by(referencia="KA")) is None
    assert db.session.scalar(db.select(Cliente).filter_by(nombre="Clínica Uno")) is None


def test_errores_de_empresa(app, tmp_path):
    vacia = [None] * 16
    filas = []
    for cambios in ({3: "T1", 9: "Marte"}, {3: "T2", 10: "D"}, {3: "T3", 12: "Quizá"},
                    {3: "T4", 14: "clia@ejemplo.com"}, {3: "T5", 4: "dos@a.es, tres@a.es"}):
        fila = list(vacia)
        fila[0] = f"Empresa {cambios[3]}"
        for i, v in cambios.items():
            fila[i] = v
        filas.append(fila)
    r = _importar(app, _excel(tmp_path, Empresas=filas))
    assert r.exit_code == 1
    for texto in ["IVA «Marte»: tiene que ser España, UE o Exportación", "«Formato de factura» tiene que ser A, B o C",
                  "«Ve precios» tiene que ser Sí o No: Quizá", "«clia@ejemplo.com» no es un usuario de Soporte de cliente",
                  "no parece válido (uno solo por empresa)"]:
        assert texto in r.output, texto


def test_archivo_que_no_es_la_plantilla(app, tmp_path):
    libro = Workbook()
    libro.active.title = "Hoja1"
    ruta = os.path.join(tmp_path, "otro.xlsx")
    libro.save(ruta)
    r = _importar(app, ruta)
    assert r.exit_code == 1 and "no tiene ninguna hoja «Productos», «Kits» ni «Empresas»" in r.output
    roto = os.path.join(tmp_path, "roto.xlsx")
    with open(roto, "w") as f:
        f.write("no soy un excel")
    assert "No se puede abrir" in _importar(app, roto).output


def test_la_plantilla_vacia_del_proyecto_tiene_todas_las_columnas():
    """deploy/plantilla_carga.xlsx: cada columna que lee el importador está en su hoja."""
    ruta = os.path.join(os.path.dirname(__file__), "..", "deploy", "plantilla_carga.xlsx")
    libro = load_workbook(ruta, read_only=True)
    for hoja, columnas in COLUMNAS.items():
        titulos = [_sin_tildes(c or "").replace("*", "").strip() for c in next(libro[hoja].iter_rows(values_only=True))]
        for campo, principios in columnas:
            assert any(t.startswith(p) for t in titulos for p in principios), (hoja, campo)
