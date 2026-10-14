"""Paso 92 · Exportar todo lo de un cliente en un ZIP (solo admin, queda en la auditoría)."""
import io
import zipfile

from openpyxl import load_workbook

from app.extensions import db
from app.models import Pedido, RegistroAuditoria
from conftest import login, pedido_cliente


def test_admin_exporta_un_zip_con_pedidos_proformas_y_documentos(app, client):
    login(client, "cliA")
    pedido_cliente(client, referencia="OC-77")
    pedido = db.session.scalar(db.select(Pedido))
    client.post(f"/pedidos/{pedido.id}/documentos",
                data={"documento": (io.BytesIO(b"%PDF-1.4 orden de compra"), "Orden: compra.pdf")},
                content_type="multipart/form-data")
    client.post("/auth/logout")
    # Una proforma antigua, sin PDF adjunto: va con el modelo que genera la plataforma
    from datetime import date
    from app.models import Proforma
    db.session.add(Proforma(numero="PF-2026-0001", pedido=pedido, valida_hasta=date(2026, 12, 31)))
    db.session.commit()
    login(client, "admin")
    assert "Exportar todo" in client.get("/clientes/1").get_data(as_text=True)
    r = client.post("/clientes/1/exportar")
    assert r.status_code == 200 and r.mimetype == "application/zip"
    assert "attachment" in r.headers["Content-Disposition"]
    z = zipfile.ZipFile(io.BytesIO(r.data))
    nombres = z.namelist()
    assert "pedidos.xlsx" in nombres and "LEEME.txt" in nombres
    proformas = [n for n in nombres if n.startswith("proformas/")]
    assert len(proformas) == 1 and z.read(proformas[0]).startswith(b"%PDF")
    documentos = [n for n in nombres if n.startswith("documentos/")]
    assert len(documentos) == 1 and documentos[0].endswith("Orden_ compra.pdf")      # (sin ":" en el nombre)
    hoja = load_workbook(io.BytesIO(z.read("pedidos.xlsx"))).active
    assert any("OC-77" == c.value for fila in hoja.iter_rows() for c in fila)
    assert "Clínica A" in z.read("LEEME.txt").decode()
    registro = db.session.scalar(db.select(RegistroAuditoria).filter_by(accion="exportar", entidad="Cliente"))
    assert registro is not None and "1 pedidos" in registro.detalle


def test_solo_admin_exporta(app, client):
    login(client, "loga")                               # Soporte de cliente de Clínica A: no
    assert "Exportar todo" not in client.get("/clientes/1").get_data(as_text=True)
    assert client.post("/clientes/1/exportar").status_code == 403
