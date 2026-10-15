"""Facturas hechas a mano: "Por facturar", "Datos para facturar" (para la plantilla de Excel)
y el registro sin importes."""
from datetime import date
from decimal import Decimal
from io import BytesIO

from app.extensions import db
from app.facturas.datos import datos_para_facturar, totales
from app.models import Factura, LineaPedido, ahora, Lote, Pedido, Producto, ReservaLote
from app.pedidos.pendientes import pendientes_de
from app.models import User
from tests.conftest import login, pedido_cliente


def _enviado(client, cantidad=2, estado="enviado"):
    """Un pedido de la Clínica A ya enviado (el camino hasta "enviado" se prueba en otros tests)."""
    login(client, "cliA"); pedido_cliente(client, cantidad=cantidad); client.post("/auth/logout")
    p = db.session.scalar(db.select(Pedido).order_by(Pedido.id.desc()))
    p.estado, p.transportista = estado, "DHL"
    db.session.commit()
    return p


def _grupo_facturar(app, email):
    """El grupo "Por facturar" de los Pendientes de ese usuario (None si no le sale)."""
    from flask_login import login_user
    usuario = db.session.scalar(db.select(User).filter_by(email=f"{email}@ejemplo.com"))
    with app.test_request_context():
        login_user(usuario)
        return next((g for g in pendientes_de(usuario, con_vacios=True) if g["clave"] == "facturar"), None)


def test_enviado_sin_factura_sale_por_facturar_hasta_que_se_registra(app, client):
    p = _enviado(client)
    login(client, "fac")
    assert p.referencia_interna in client.get("/facturas/por-facturar").get_data(as_text=True)
    assert [f["pedido"].id for f in _grupo_facturar(app, "fac")["todas"]] == [p.id]
    r = client.post(f"/facturas/pedido/{p.id}", content_type="multipart/form-data", data={
        "numero": "F-2026-0145", "fecha_emision": "2026-10-07", "fecha_vencimiento": "2026-11-06",
        "pdf": (BytesIO(b"%PDF-1.4\n" + b"0" * 300), "F-2026-0145.pdf")})
    assert r.status_code == 302
    f = db.session.scalar(db.select(Factura))
    assert f.numero == "F-2026-0145" and f.total is None and f.pdf_archivo
    assert p.referencia_interna not in client.get("/facturas/por-facturar").get_data(as_text=True)
    # Si se anula, el pedido vuelve a estar por facturar
    f.estado = "anulada"; db.session.commit()
    assert p.referencia_interna in client.get("/facturas/por-facturar").get_data(as_text=True)


def test_por_facturar_y_datos_solo_para_quien_registra_facturas(app, client):
    p = _enviado(client)
    for quien in ("loga", "cliA"):               # el FAS del cliente y el propio cliente: no
        login(client, quien)
        assert client.get("/facturas/por-facturar").status_code == 403
        assert client.get(f"/facturas/pedido/{p.id}/datos").status_code == 403
        client.post("/auth/logout")
    assert _grupo_facturar(app, "loga") is None


def _excel(client, p, consulta=""):
    """Descarga el Excel de "Preparar factura" y lo devuelve como {celda: valor} con las fórmulas
    ya calculadas (las pocas que usa: SUM, ROUND, +, *, /), y la hoja para mirar el resto."""
    import re
    from io import BytesIO
    from openpyxl import load_workbook
    r = client.get(f"/facturas/pedido/{p.id}/datos.xlsx{consulta}")
    assert r.status_code == 200 and r.headers["Content-Disposition"].endswith(f'factura_{p.referencia_interna}.xlsx"')
    ws = load_workbook(BytesIO(r.data)).active
    crudo = {c.coordinate: c.value for fila in ws.iter_rows() for c in fila if c.value is not None}

    def valor(ref):
        v = crudo.get(ref, 0)
        if not (isinstance(v, str) and v.startswith("=")):
            return v
        expr = re.sub(r"SUM\(([A-Z]+)(\d+):([A-Z]+)(\d+)\)",
                      lambda m: "(" + "+".join(f"{m.group(1)}{n}" for n in range(int(m.group(2)), int(m.group(4)) + 1)) + ")",
                      v[1:]).replace("ROUND", "round")
        return round(eval(re.sub(r"\b([A-Z]+\d+)\b", lambda m: repr(valor(m.group(1))), expr)), 2)
    return {ref: valor(ref) for ref in crudo}, ws


def _totales(celdas):
    """{nombre del total: importe} leyendo las filas "Algo:" del Excel."""
    from openpyxl.utils import column_index_from_string, get_column_letter
    res = {}
    for ref, v in celdas.items():
        if isinstance(v, str) and v.endswith(":"):
            col = "".join(ch for ch in ref if ch.isalpha()); fila = ref[len(col):]
            importe = celdas.get(f"{get_column_letter(column_index_from_string(col) + 1)}{fila}")
            if importe is not None:                  # (los textos de la transferencia también acaban en «:»)
                res[v[:-1]] = importe
    return res


def test_factura_en_excel_formato_a_por_defecto(app, client):
    p = _enviado(client, cantidad=2)
    p1 = db.session.scalar(db.select(Producto).filter_by(referencia="P1"))
    lote = Lote(producto=p1, codigo="YK2501-C3", caducidad=date(2027, 11, 30))
    db.session.add(ReservaLote(pedido=p, lote=lote, cantidad=2)); db.session.commit()
    login(client, "fac")
    html = client.get(f"/facturas/pedido/{p.id}/datos").get_data(as_text=True)
    assert "YK2501-C3" in html and "30/11/2027" in html                     # el lote se ve en la página...
    assert "no tiene configurada la facturación" in html                    # ...y avisa: cliente sin configurar
    celdas, ws = _excel(client, p)
    textos = [v for v in celdas.values() if isinstance(v, str)]
    assert "YK2501-C3" not in textos                                         # ...pero no va en el Excel
    # La factura completa: Seller / Billing / Delivery (sin datos de envío, la de facturación), nº y fecha
    assert {"Seller.", "Billing Address.", "Delivery address", "INVOICE"} <= set(textos)
    assert textos.count("Clínica A") == 2 and textos.count("VAT: B11111111") == 2
    assert f"Invoice Number: {p.referencia_interna}" in textos
    todo = " ".join(textos)
    assert "SWIFT CODE: CAIXESBBXXX" in todo and "Shipped by: DHL" in todo and "EORI" not in todo   # (España: sin EORI)
    assert len(ws._images) == 3                                              # logo, transferencia y sello
    # Formato A: PRODUCT (A:C combinada), Unit, Catalog No., Item, Qty., HS Code; sin precios por línea
    fila = next(c.row for c in ws["A"] if c.value == "PRODUCT")
    assert [ws.cell(fila, k).value for k in (1, 4, 5, 6, 7, 8)] == ["PRODUCT", "Unit", "Catalog No.", "Item", "Qty.", "HS Code"]
    assert [ws.cell(fila + 1, k).value for k in (1, 4, 5, 6, 7, 8)] == ["Kit uno", 2, "P1", "Kit 1", 2, 3822190090]
    # Sin configurar: IVA España
    assert _totales(celdas) == {"Total price": 200, "% IVA": 21, "Total tax": 42, "Total amount": 242}


def test_descuentos_por_linea_dan_el_porcentaje_equivalente(app, client):
    p = _enviado(client, cantidad=2)              # 2 kits a 100 €
    p1 = db.session.scalar(db.select(Producto).filter_by(referencia="P1"))
    # Un suelto de 200 € con un 50 % de descuento: 400 € sin descuento, 300 € con él → 25 %
    p.lineas.append(LineaPedido(producto=p1, cantidad=2, precio_unitario=100, iva_pct=21, descuento_pct=50))
    db.session.commit()
    t = totales(p)
    assert (t["bruto"], t["neto"], t["descuento"]) == (Decimal("400.00"), Decimal("300.00"), Decimal("25.00"))
    assert t["iva_pct"] == 21 and t["total"] == Decimal("363.00")


def test_la_lista_de_facturas_pone_primero_las_vencidas(app, client):
    p = _enviado(client)
    db.session.add_all([
        Factura(enviada_el=ahora(), numero="F-NUEVA", cliente=p.cliente, fecha_emision=date(2026, 10, 1), fecha_vencimiento=date(2099, 1, 1)),
        Factura(enviada_el=ahora(), numero="F-VENCIDA", cliente=p.cliente, fecha_emision=date(2026, 1, 1), fecha_vencimiento=date(2026, 2, 1))])
    db.session.commit()
    login(client, "fac")
    html = client.get("/facturas/").get_data(as_text=True)
    assert html.index("F-VENCIDA") < html.index("F-NUEVA")


def _configurar(client, p, formato, regimen, envio=""):
    return client.post(f"/facturas/cliente/{p.cliente_id}/facturacion", data={
        "formato_factura": formato, "regimen_iva": regimen, "gastos_envio": envio, "pedido_id": p.id})


def test_formato_b_ue_con_envio_pactado(app, client):
    p = _enviado(client, cantidad=2)                     # 2 × "Kit uno" a 100 € = 2 × P1 (catálogo 100 €)
    login(client, "fac")
    r = _configurar(client, p, "B", "ue", "150")
    assert r.status_code == 302 and f"/facturas/pedido/{p.id}/datos" in r.headers["Location"]
    assert datos_para_facturar(p)["avisos"] == []                                  # 2 × 100 = el precio del kit
    celdas, ws = _excel(client, p)
    fila = next(c.row for c in ws["A"] if c.value == "PRODUCT")
    assert ws.cell(fila, 9).value == "Price" and ws.cell(fila + 1, 9).value == 200     # precio de la línea
    # UE: sin IVA, con el envío pactado
    assert _totales(celdas) == {"Total price": 200, "Shipment": 150, "Total amount": 350}
    # Otro envío solo para esta factura: cambia la página y el Excel, no el cliente
    html = client.get(f"/facturas/pedido/{p.id}/datos?envio=30").get_data(as_text=True)
    assert "230,00" in html and "Envío cambiado solo aquí" in html and "datos.xlsx?envio=30" in html
    assert _totales(_excel(client, p, "?envio=30")[0])["Total amount"] == 230
    assert p.cliente.gastos_envio == Decimal("150.00")


def test_formato_b_avisa_si_los_productos_no_suman_el_kit(app, client):
    p = _enviado(client, cantidad=2)
    p.kits[0].precio_unitario = 90                        # el kit a 90 €; sus productos, a 100 € de catálogo
    p.cliente.formato_factura, p.cliente.regimen_iva = "B", "es"
    db.session.commit()
    assert datos_para_facturar(p)["avisos"] == [("Kit uno", Decimal("200.00"), Decimal("180.00"))]


def test_formato_c_lista_con_precio_unitario_y_descuento(app, client):
    p = _enviado(client, cantidad=1)
    p1 = db.session.scalar(db.select(Producto).filter_by(referencia="P1"))
    p.lineas.append(LineaPedido(producto=p1, cantidad=3, precio_unitario=100, iva_pct=21, descuento_pct=50))
    p.cliente.formato_factura, p.cliente.regimen_iva = "C", "export"
    db.session.commit()
    login(client, "fac")
    celdas, ws = _excel(client, p)
    fila = next(c.row for c in ws["A"] if c.value == "Item")
    # Item (A:C), Catalog No., HS Code, Unit Price, Qty., Price (de la línea, sin descuento: va en los totales)
    cols = (1, 4, 5, 6, 7, 8)
    assert [ws.cell(fila, k).value for k in cols] == ["Item", "Catalog No.", "HS Code", "Unit Price", "Qty.", "Total Price"]
    assert [[ws.cell(fila + n, k).value for k in cols] for n in (1, 2)] == [
        ["Kit 1", "P1", 3822190090, 100, 1, 100], ["Kit 1", "P1", 3822190090, 100, 3, 300]]
    assert _totales(celdas) == {"Total before discount": 400, "Discount [%]": 37.5,
                                "Subtotal less discount": 250, "Total amount": 250}       # exportación: sin IVA


def test_en_espana_el_envio_lleva_iva(app, client):
    p = _enviado(client, cantidad=1)
    t = totales(p, "es", 100)
    assert (t["iva"], t["total"]) == (Decimal("42.00"), Decimal("242.00"))     # 21 % de 100 + 21 % de 100


def test_facturacion_del_cliente_solo_la_cambia_quien_registra_facturas(app, client):
    p = _enviado(client)
    login(client, "loga")                                  # FAS: no
    assert _configurar(client, p, "C", "ue").status_code == 403
    client.post("/auth/logout"); login(client, "fac")
    _configurar(client, p, "Z", "ue")                      # formato que no existe: no se guarda
    assert p.cliente.formato_factura is None


def test_el_numero_de_factura_viene_puesto_con_la_referencia(app, client):
    p = _enviado(client)
    login(client, "fac")
    html = client.get(f"/pedidos/{p.id}").get_data(as_text=True)
    import re
    campo = re.search(r'<input[^>]*name="numero"[^>]*>', html).group(0)
    assert f'value="{p.referencia_interna}"' in campo


# --------------------------------------------------------------------------- #
# Revisión por Soporte de cliente y envío al cliente
# --------------------------------------------------------------------------- #
def _registrar_factura(client, p, revisar):
    datos = {"numero": p.referencia_interna, "fecha_emision": "2026-10-08", "fecha_vencimiento": "2026-11-07",
             "pdf": (BytesIO(b"%PDF-1.4\n" + b"0" * 300), "f.pdf")}
    if revisar:
        datos["revisar"] = "y"
    client.post(f"/facturas/pedido/{p.id}", data=datos, content_type="multipart/form-data")
    return db.session.scalar(db.select(Factura).filter_by(pedido_id=p.id))


def _cliente_la_ve(client, f):
    client.post("/auth/logout"); login(client, "cliA")
    ve = f.numero in client.get("/facturas/mias").get_data(as_text=True)
    ve_pdf = client.get(f"/facturas/{f.id}/pdf").status_code == 200
    assert ve == ve_pdf
    client.post("/auth/logout")
    return ve


def test_la_casilla_de_revision_sale_marcada_si_el_cliente_lo_tiene(app, client):
    p = _enviado(client)
    login(client, "fac")
    _configurar(client, p, "A", "es")                      # sin la casilla: revisión desactivada
    assert not p.cliente.revisar_facturas
    client.post(f"/facturas/cliente/{p.cliente_id}/facturacion", data={
        "formato_factura": "A", "regimen_iva": "es", "revisar_facturas": "1"})
    assert p.cliente.revisar_facturas
    import re
    html = client.get(f"/pedidos/{p.id}").get_data(as_text=True)
    assert "checked" in re.search(r'<input[^>]*name="revisar"[^>]*>', html).group(0)


def test_revision_devolver_corregir_aprobar_y_enviar(app, client):
    p = _enviado(client)
    login(client, "fac")
    f = _registrar_factura(client, p, revisar=True)
    assert f.fase == "revision" and not _cliente_la_ve(client, f)
    # La revisa el Soporte de cliente de ese cliente; otro no puede
    login(client, "logb")
    assert client.post(f"/facturas/{f.id}/aprobar").status_code == 404
    client.post("/auth/logout"); login(client, "loga")
    assert [x["texto"] for x in _grupo(app, "loga", "facturas_revisar")["todas"]] == [f.numero]
    client.post(f"/facturas/{f.id}/devolver", data={"comentario": ""})          # sin comentario: no
    assert f.fase == "revision"
    client.post(f"/facturas/{f.id}/devolver", data={"comentario": "Falta el descuento del 10 %"})
    assert f.fase == "cambios" and f.revision_comentario == "Falta el descuento del 10 %"
    # Aina la ve devuelta, sube el PDF corregido y vuelve a revisión
    client.post("/auth/logout"); login(client, "fac")
    assert [x["texto"] for x in _grupo(app, "fac", "facturas_devueltas")["todas"]] == [f.numero]
    assert client.post(f"/facturas/{f.id}/enviada").status_code == 302 and not f.enviada   # devuelta: no se envía
    client.post(f"/facturas/{f.id}/subir-pdf", data={"pdf": (BytesIO(b"%PDF-1.4\n" + b"1" * 300), "f2.pdf")},
                content_type="multipart/form-data")
    assert f.fase == "revision"
    client.post("/auth/logout"); login(client, "loga")
    client.post(f"/facturas/{f.id}/aprobar")
    assert f.fase == "por_enviar" and f.revisado_por.email == "loga@ejemplo.com"
    assert not _cliente_la_ve(client, f)                   # aprobada, pero aún no enviada
    login(client, "fac")
    assert [x["texto"] for x in _grupo(app, "fac", "facturas_enviar")["todas"]] == [f.numero]
    client.post(f"/facturas/{f.id}/enviada")
    assert f.enviada and _cliente_la_ve(client, f)


def test_sin_revision_va_directa_a_por_enviar(app, client):
    p = _enviado(client)
    login(client, "fac")
    f = _registrar_factura(client, p, revisar=False)
    assert f.fase == "por_enviar" and f.revision is None
    f.fecha_vencimiento = date(2020, 1, 1); db.session.commit()
    assert not f.esta_vencida                              # sin enviar no cuenta como vencida
    client.post(f"/facturas/{f.id}/enviada")
    assert f.esta_vencida and _cliente_la_ve(client, f)


def _grupo(app, email, clave):
    from flask_login import login_user
    usuario = db.session.scalar(db.select(User).filter_by(email=f"{email}@ejemplo.com"))
    with app.test_request_context():
        login_user(usuario)
        return next((g for g in pendientes_de(usuario, con_vacios=True) if g["clave"] == clave), None)


def test_factura_de_un_cliente_de_la_ue_lleva_los_eori(app, client):
    """15/10/2026 · Cliente de la UE: el EORI de Yikon en Seller y el suyo en su dirección. Fuera de la
    UE, ninguno. La dirección de entrega, la de sus datos de envío; el uso y su nº de pedido."""
    from app import datos_empresa
    p = _enviado(client, cantidad=1)
    c = p.cliente
    c.pais, c.eori, c.ciudad, c.direccion = "DE", "DE123456789", "Berlin", "Hauptstr. 1"
    c.envio_igual, c.envio_direccion, c.envio_ciudad, c.envio_pais = False, "Lab Str. 9", "Hamburg", "DE"
    p.referencia, p.uso = "PO-77", "ivd"
    datos_empresa.guardar({**datos_empresa.datos(), "eori": "ESB44777787", "incoterm": "DAP"})
    db.session.commit()
    login(client, "fac")
    celdas, ws = _excel(client, p)
    textos = [v for v in celdas.values() if isinstance(v, str)]
    todo = " | ".join(textos)
    assert "EORI: ESB44777787" in textos and "EORI: DE123456789" in textos
    assert "GERMANY" in textos and "Lab Str. 9" in textos and "Hamburg" in " ".join(textos)
    assert "Your PO: PO-77" in textos and "USE: FOR IN VITRO DIAGNOSTIC USE (IVD)" in textos
    assert "6. Incoterm: DAP" in todo and "REFERENCE CONTRACT NUMBER: PO-77" in todo
    # Fuera de la UE: sin EORI
    c.pais = "MX"
    db.session.commit()
    textos = [v for v in _excel(client, p)[0].values() if isinstance(v, str)]
    assert not any(t.startswith("EORI") for t in textos)


def test_el_hs_code_es_de_cada_producto(app, client):
    p = _enviado(client, cantidad=1)
    login(client, "admin")
    p1 = db.session.scalar(db.select(Producto).filter_by(referencia="P1"))
    assert p1.codigo_hs == "3822190090"                            # el de la plantilla, por defecto
    assert 'name="codigo_hs"' in client.get(f"/productos/{p1.id}/editar").get_data(as_text=True)
    p1.codigo_hs = "3002900000"
    db.session.commit()
    assert datos_para_facturar(p)["filas"][0]["hs"] == "3002900000"
