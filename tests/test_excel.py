"""Exportación a Excel: respeta filtros, alcance y permisos (el cliente, sin importes)."""
from io import BytesIO

from openpyxl import load_workbook

from app.models import ahora  # facturas ya enviadas al cliente (las ve)
from app.extensions import db
from app.models import Factura, Pedido
from tests.conftest import login, pedido_cliente


def _leer(respuesta):
    assert respuesta.status_code == 200
    assert respuesta.mimetype.endswith("spreadsheetml.sheet")
    ws = load_workbook(BytesIO(respuesta.data)).active
    return [[c.value for c in fila] for fila in ws.iter_rows()]


def test_excel_de_pedidos_con_importes_para_admin(app, client):
    login(client, "cliA"); pedido_cliente(client, referencia="OC-XL", cantidad=2); client.post("/auth/logout")
    login(client, "admin")
    filas = _leer(client.get("/pedidos/exportar.xlsx"))
    cabecera = filas[2]                                  # fila 1: filtros · fila 3: cabecera
    assert "Total" in cabecera and "Productos" in cabecera
    datos = filas[3]
    assert "OC-XL" in datos and "2 × Kit uno (2 × Kit 1)" in datos and 242.0 in datos   # 2 kits × 100 € + 21 % IVA


def test_excel_del_cliente_sin_importes_y_solo_lo_suyo(app, client):
    login(client, "cliB"); pedido_cliente(client, referencia="DE-B"); client.post("/auth/logout")
    login(client, "cliA"); pedido_cliente(client, referencia="DE-A")
    filas = _leer(client.get("/pedidos/exportar.xlsx"))
    todo = str(filas)
    assert "DE-A" in todo and "DE-B" not in todo
    assert "Total" not in filas[2] and "IVA" not in filas[2]


def test_excel_de_facturas(app, client):
    from datetime import date
    db.session.add(Factura(enviada_el=ahora(), numero="F-77", cliente_id=1, fecha_emision=date(2026, 9, 1),
                           base_imponible=100, iva=21, total=121))
    db.session.commit()
    login(client, "admin")
    filas = _leer(client.get("/facturas/exportar.xlsx"))
    assert filas[0][0] == "Número" and filas[1][0] == "F-77" and "Total" not in filas[0]   # sin importes
    # FAS: solo las facturas de sus clientes (F-77 es de la Clínica A, que lleva "loga")
    client.post("/auth/logout"); login(client, "logb")
    filas = _leer(client.get("/facturas/exportar.xlsx"))
    assert all(f[0] != "F-77" for f in filas)
    client.post("/auth/logout"); login(client, "loga")
    assert any(f[0] == "F-77" for f in _leer(client.get("/facturas/exportar.xlsx")))
